# -*- coding: utf-8 -*-
"""Importación puntual: vuelca 'cartera en excel.xlsx' (hojas de
valores y operaciones) en la base de datos apuntada por DATABASE_URL
(ver rxconfig.py) -- pensado para ejecutarse contra el Postgres de
Railway, con este archivo en la raíz del proyecto.

Ejecutar con (openpyxl no es dependencia del proyecto, se instala solo
para esta ejecución):
    uv run --with openpyxl python importar_cartera_excel.py

Antes de cada cartera nueva, cambia EMAIL / TIPO_CARTERA / BROKER más
abajo y sustituye 'cartera en excel.xlsx' por el Excel de esa cartera.

Es idempotente para VALORES (si el ticker+mercado ya existe, reutiliza
ese id_valor en vez de duplicarlo) pero NO para OPERACIONES: por eso
comprueba antes si ya hay operaciones para esta cartera+bróker y para
en seco si las hay, en vez de asumir que está vacío.
"""

import sys

import openpyxl

from gestion_cartera.auth_db import obtener_usuario_por_email
from gestion_cartera.operaciones_db import (
    crear_operacion,
    crear_valor,
    obtener_broker_id_por_nombre,
    obtener_cartera_id,
    obtener_operaciones,
    obtener_sector_id,
    obtener_valor_id_por_ticker_mercado,
)

EXCEL_PATH = "cartera en excel.xlsx"
EMAIL = "marcgarrido93@gmail.com"
TIPO_CARTERA = "Largo Plazo"
BROKER = "ING"

MONEDA_POR_MERCADO = {
    "BME": "EUR",
    "EPA": "EUR",
    "AMS": "EUR",
    "ETR": "EUR",
    "LON": "GBP",
    "NASDAQ": "USD",
    "NYSE": "USD",
}

# "MERCADO:TICKER" (tal como aparece en la hoja de valores) -> (supersector,
# sector, grupo) de Morningstar. Acumula el mapeo de todas las carteras ya
# importadas, para no tener que repetir trabajo con valores repetidos.
MAPEO_SECTORES = {
    # --- cartera gasbis (ING, Largo Plazo) ---
    "LON:IDS": ("Sensible", "Industria", "Transporte"),
    "NYSE:KO": ("Defensivo", "Consumo Defensivo", "Bebidas No Alcohólicas"),
    "BME:ELE": ("Defensivo", "Servicios Públicos", "Eléctricas Reguladas"),
    "EPA:BN": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "AMS:UNA": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "BME:ZOT": ("Sensible", "Industria", "Productos Industriales"),
    "BME:CIE": ("Cíclico", "Consumo Cíclico", "Vehículos y Componentes"),
    "BME:ENG": ("Sensible", "Energía", "Otras Fuentes de Energía"),
    "BME:ITX": ("Cíclico", "Consumo Cíclico", "Distribución - Cíclica"),
    "BME:ACX": ("Cíclico", "Materiales Básicos", "Acero"),
    "BME:LOG": ("Sensible", "Industria", "Distribución Industrial"),
    "BME:MAP": ("Cíclico", "Servicios Financieros", "Seguros"),
    "BME:CABK": ("Cíclico", "Servicios Financieros", "Banca"),
    "BME:ACS": ("Sensible", "Industria", "Construcción"),
    "BME:RED": ("Defensivo", "Servicios Públicos", "Eléctricas Reguladas"),
    "BME:BME": ("Cíclico", "Servicios Financieros", "Mercados de Capitales"),
    "BME:BBVA": ("Cíclico", "Servicios Financieros", "Banca"),
    "BME:TEF": ("Sensible", "Comunicación", "Telecomunicaciones"),
    "BME:REP": ("Sensible", "Energía", "Petróleo y Gas"),
    "NYSE:T": ("Sensible", "Comunicación", "Telecomunicaciones"),
    "NYSE:JNJ": ("Defensivo", "Sanidad", "Farmacéuticas"),
    "NASDAQ:KHC": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "NASDAQ:KMB": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "BME:AENA": ("Sensible", "Industria", "Transporte"),
    "ETR:BAS": ("Cíclico", "Materiales Básicos", "Química"),
    "BME:COL": ("Cíclico", "Inmobiliario", "REITs (Fondos de Inversión Inmobiliaria)"),
    "NASDAQ:TROW": ("Cíclico", "Servicios Financieros", "Gestión de Activos"),
    "BME:VIS": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "LON:SDR": ("Cíclico", "Servicios Financieros", "Gestión de Activos"),
    "NYSE:GIS": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "AMS:MICC": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "NASDAQ:PEP": ("Defensivo", "Consumo Defensivo", "Bebidas No Alcohólicas"),
    # --- cartera mjcaroli (ING, Largo Plazo) -- nuevos ---
    "NYSE:SOLV": ("Defensivo", "Sanidad", "Dispositivos e Instrumental Médico"),
    "NYSE:VZ": ("Sensible", "Comunicación", "Telecomunicaciones"),
    "LON:NG": ("Defensivo", "Servicios Públicos", "Eléctricas Reguladas"),
    "NYSE:KHC": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "LON:BT.A": ("Sensible", "Comunicación", "Telecomunicaciones"),
    "BME:LDA": ("Cíclico", "Servicios Financieros", "Seguros"),
    "LON:ABF": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "NYSE:MMM": ("Sensible", "Industria", "Productos Industriales"),
    "ETR:HEN": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "ETR:BAYN": ("Defensivo", "Sanidad", "Farmacéuticas"),
    "NYSE:SPG": ("Cíclico", "Inmobiliario", "REITs (Fondos de Inversión Inmobiliaria)"),
    "NASDAQ:INTC": ("Sensible", "Tecnología", "Semiconductores"),
    "BME:MEL": ("Cíclico", "Consumo Cíclico", "Viajes y Ocio"),
    "BME:MRL": ("Cíclico", "Inmobiliario", "REITs (Fondos de Inversión Inmobiliaria)"),
    "BME:ENC": ("Cíclico", "Materiales Básicos", "Productos Forestales"),
    "BME:IBE": ("Defensivo", "Servicios Públicos", "Eléctricas Reguladas"),
    "BME:SAN": ("Cíclico", "Servicios Financieros", "Banca"),
    "BME:ANA": ("Sensible", "Industria", "Construcción"),
    # --- cartera mjcaroli (ING, Corto Plazo) -- nuevos ---
    "NASDAQ:SPCX": ("Sensible", "Industria", "Aeroespacial y Defensa"),
    "NASDAQ:GFS": ("Sensible", "Tecnología", "Semiconductores"),
    "BME:COM": ("Sensible", "Industria", "Servicios Empresariales"),
    "BME:EZE": ("Sensible", "Industria", "Servicios Empresariales"),
    "BME:ART": ("Sensible", "Industria", "Productos Industriales"),
    "BME:ZNK": ("Sensible", "Comunicación", "Medios de Comunicación Diversificados"),
    "BME:MTS": ("Cíclico", "Materiales Básicos", "Acero"),
    "BME:SCYR": ("Sensible", "Industria", "Construcción"),
    # --- cartera joangc93 (ING, Largo Plazo) -- nuevos ---
    "LON:AV": ("Cíclico", "Servicios Financieros", "Seguros"),
    "BME:HOME": ("Cíclico", "Consumo Cíclico", "Construcción de Viviendas"),
    "BME:EBRO": ("Defensivo", "Consumo Defensivo", "Bienes de Consumo Envasados"),
    "NYSE:CVX": ("Sensible", "Energía", "Petróleo y Gas"),
    "BME:SAB": ("Cíclico", "Servicios Financieros", "Banca"),
    # --- cartera marcgarrido93 (ING, Largo Plazo) -- nuevos ---
    "EPA:MC": ("Cíclico", "Consumo Cíclico", "Fabricación - Textil y Mobiliario"),
    "ETR:BMW": ("Cíclico", "Consumo Cíclico", "Vehículos y Componentes"),
    "NYSE:TAP": ("Defensivo", "Consumo Defensivo", "Bebidas Alcohólicas"),
    "NASDAQ:CMCSA": ("Sensible", "Comunicación", "Medios de Comunicación Diversificados"),
    "LON:DGE": ("Defensivo", "Consumo Defensivo", "Bebidas Alcohólicas"),
}

TIPO_OPERACION = {
    "COMPRA": "Compra",
    "VENTA": "Venta",
    "DIVIDENDO": "Dividendo",
    "PRIMA": "Prima",
    "SCRIPT": "Script",
    "SCR_CPA": "Script",
}
TIPO_DERECHO_SCRIPT = {
    "SCRIPT": "Venta",  # derecho vendido
    "SCR_CPA": "Compra",  # derecho comprado
}


def _num(v):
    return float(v) if v is not None else None


def _hoja(wb, *nombres_posibles):
    """Busca una hoja por nombre, sin importar mayúsculas/minúsculas."""
    por_minusculas = {n.lower(): n for n in wb.sheetnames}
    for nombre in nombres_posibles:
        real = por_minusculas.get(nombre.lower())
        if real:
            return wb[real]
    sys.exit(f"No encuentro ninguna hoja llamada {nombres_posibles} en {EXCEL_PATH}")


def main() -> None:
    usuario = obtener_usuario_por_email(EMAIL)
    if usuario is None:
        sys.exit(f"No existe el usuario {EMAIL}")

    id_cartera = obtener_cartera_id(usuario.id, TIPO_CARTERA)
    if id_cartera is None:
        sys.exit(f"No existe la cartera «{TIPO_CARTERA}» para {EMAIL}")

    id_broker = obtener_broker_id_por_nombre(BROKER)
    if id_broker is None:
        sys.exit(f"No existe el bróker «{BROKER}»")

    ya_existentes = obtener_operaciones(id_cartera)
    ya_de_este_broker = [o for o in ya_existentes if o.get("broker") == BROKER]
    if ya_de_este_broker:
        sys.exit(
            f"Ya hay {len(ya_de_este_broker)} operaciones en «{TIPO_CARTERA}» "
            f"de {EMAIL} con bróker «{BROKER}». Aborto para no duplicar -- si "
            "quieres reimportar, bórralas primero desde la app."
        )

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    # --- 1. Valores ---
    ws_valores = _hoja(wb, "Valores", "valores")
    headers = [c.value for c in ws_valores[1]]
    id_valor_por_ticker: dict[str, int] = {}
    creados = existentes = 0
    for row in ws_valores.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        ticker_completo = d.get("Ticker")
        if not ticker_completo:
            continue
        mercado, ticker = ticker_completo.split(":", 1)
        zona = d.get("Zona")
        empresa = d.get("Empresa")
        moneda = MONEDA_POR_MERCADO[mercado]

        id_existente = obtener_valor_id_por_ticker_mercado(ticker, mercado)
        if id_existente is not None:
            id_valor_por_ticker[ticker_completo] = id_existente
            existentes += 1
            continue

        mapeo = MAPEO_SECTORES.get(ticker_completo)
        if mapeo is None:
            sys.exit(
                f"Falta el mapeo de sector para {ticker_completo} "
                "(añádelo a MAPEO_SECTORES antes de reintentar)."
            )
        supersector, sector, grupo = mapeo
        id_sector = obtener_sector_id(supersector, sector, grupo)
        if id_sector is None:
            sys.exit(
                f"No se encuentra el sector {supersector}/{sector}/{grupo} "
                f"(valor {ticker_completo}) -- revisa nombres exactos."
            )

        id_valor = crear_valor(ticker, empresa, id_sector, mercado, zona, moneda)
        id_valor_por_ticker[ticker_completo] = id_valor
        creados += 1

    print(f"Valores: {creados} creados, {existentes} ya existían.")

    # --- 2. Operaciones ---
    ws_ops = _hoja(wb, "operaciones", "Operaciones")
    headers_ops = [c.value for c in ws_ops[1]]
    creadas = 0
    for row in ws_ops.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers_ops, row))
        operac = d.get("Operac")
        if not operac:
            continue
        ticker_completo = d.get("Ticker")
        id_valor = id_valor_por_ticker.get(ticker_completo)
        if id_valor is None:
            sys.exit(f"Operación con ticker desconocido: {ticker_completo}")

        tipo_operacion = TIPO_OPERACION[operac]
        tipo_derecho = TIPO_DERECHO_SCRIPT.get(operac)

        fecha_raw = d.get("Fecha")
        fecha = fecha_raw.date() if hasattr(fecha_raw, "date") else fecha_raw

        num_titulos = _num(d.get("NumTits")) or 0.0
        importe = _num(d.get("Importe")) or 0.0
        importe_unitario = None if tipo_operacion == "Script" else _num(d.get("ImpUnit"))
        retencion_origen = _num(d.get("RetOrig"))
        retencion_destino = _num(d.get("RetDest"))
        observaciones = d.get("Observaciones")

        crear_operacion(
            id_cartera=id_cartera,
            id_valor=id_valor,
            id_broker=id_broker,
            tipo_operacion=tipo_operacion,
            fecha=fecha,
            num_titulos=num_titulos,
            importe=importe,
            importe_unitario=importe_unitario,
            retencion_origen=retencion_origen,
            retencion_destino=retencion_destino,
            tipo_derecho_script=tipo_derecho,
            observaciones=observaciones,
        )
        creadas += 1

    print(f"Operaciones creadas: {creadas}")


if __name__ == "__main__":
    main()
