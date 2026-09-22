"""Genera el Excel descargable de la página IRPF (tarea 2): mismo
resumen que se ve en pantalla (ver resumen_irpf_db.obtener_resumen_irpf),
en una única hoja, con los importes como números reales (no texto), para
que se puedan sumar/ordenar/filtrar directamente en el Excel.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gestion_cartera.resumen_irpf_db import obtener_resumen_irpf

FUENTE = "Arial"
FORMATO_EUR = '#,##0.00" €"'
FORMATO_TITULOS = "#,##0.####"

_FONT_TITULO = Font(name=FUENTE, size=14, bold=True)
_FONT_SECCION = Font(name=FUENTE, size=12, bold=True)
_FONT_CABECERA = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
_FILL_CABECERA = PatternFill("solid", fgColor="153547")  # var(--accent-3) de la app
_FONT_NORMAL = Font(name=FUENTE, size=10)
_FONT_TOTAL = Font(name=FUENTE, size=10, bold=True)
_FONT_NOTA = Font(name=FUENTE, size=9, italic=True, color="666666")


def _escribir_cabecera(ws, fila: int, columnas: list[str]) -> None:
    for col, texto in enumerate(columnas, start=1):
        celda = ws.cell(row=fila, column=col, value=texto)
        celda.font = _FONT_CABECERA
        celda.fill = _FILL_CABECERA
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autoajustar_columnas(ws, anchos: list[int]) -> None:
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def generar_excel_irpf(id_usuario: int, anio: int) -> bytes:
    resumen = obtener_resumen_irpf(id_usuario, anio)

    wb = Workbook()
    ws = wb.active
    ws.title = f"IRPF {anio}"

    fila = 1
    ws.cell(row=fila, column=1, value=f"Resumen IRPF — Año {anio}").font = _FONT_TITULO
    fila += 1
    ws.cell(
        row=fila, column=1, value="Combina Cartera de largo plazo + Cartera de corto plazo."
    ).font = _FONT_NOTA
    fila += 2

    # --- Totales -----------------------------------------------------------
    ws.cell(row=fila, column=1, value="Totales").font = _FONT_SECCION
    fila += 1
    totales = [
        ("Dividendos + venta de derechos (bruto)", resumen["total_bruto_dividendos"]),
        ("Retención en destino", resumen["total_retencion_destino"]),
        ("Retención en origen (hasta 15%)", resumen["total_retencion_origen_hasta15"]),
        ("Retención en origen (exceso sobre 15%)", resumen["total_retencion_origen_exceso"]),
        ("Plusvalía de ventas (FIFO)", resumen["total_plusvalia_ventas"]),
    ]
    for etiqueta, valor in totales:
        ws.cell(row=fila, column=1, value=etiqueta).font = _FONT_NORMAL
        celda_valor = ws.cell(row=fila, column=2, value=valor)
        celda_valor.font = _FONT_TOTAL
        celda_valor.number_format = FORMATO_EUR
        fila += 1
    fila += 1

    # --- Dividendos y venta de derechos -------------------------------------
    ws.cell(row=fila, column=1, value="Dividendos y venta de derechos").font = _FONT_SECCION
    fila += 1
    cabecera_dividendos = [
        "Fecha", "Cartera", "Ticker", "Empresa", "Tipo",
        "Bruto", "Ret. destino", "Ret. origen ≤15%", "Ret. origen >15%",
    ]
    _escribir_cabecera(ws, fila, cabecera_dividendos)
    fila_inicio_dividendos = fila + 1
    fila += 1
    for item in resumen["listado_dividendos"]:
        ws.cell(row=fila, column=1, value=item["fecha_mostrar"]).font = _FONT_NORMAL
        ws.cell(row=fila, column=2, value=item["cartera"]).font = _FONT_NORMAL
        ws.cell(row=fila, column=3, value=item["ticker"]).font = _FONT_NORMAL
        ws.cell(row=fila, column=4, value=item["empresa"]).font = _FONT_NORMAL
        ws.cell(row=fila, column=5, value=item["tipo"]).font = _FONT_NORMAL
        for col, clave in ((6, "importe"), (7, "retencion_destino"), (8, "retencion_origen_hasta15"), (9, "retencion_origen_exceso")):
            celda = ws.cell(row=fila, column=col, value=item[clave])
            celda.font = _FONT_NORMAL
            celda.number_format = FORMATO_EUR
        fila += 1
    fila_fin_dividendos = fila - 1
    if fila_fin_dividendos >= fila_inicio_dividendos:
        ws.cell(row=fila, column=5, value="Total").font = _FONT_TOTAL
        for col, letra in ((6, "F"), (7, "G"), (8, "H"), (9, "I")):
            celda = ws.cell(
                row=fila, column=col,
                value=f"=SUM({letra}{fila_inicio_dividendos}:{letra}{fila_fin_dividendos})",
            )
            celda.font = _FONT_TOTAL
            celda.number_format = FORMATO_EUR
        fila += 1
    else:
        ws.cell(row=fila, column=1, value="(sin dividendos ni venta de derechos este año)").font = _FONT_NOTA
        fila += 1
    fila += 1

    # --- Exceso de retención en origen por zona / mercado -------------------
    ws.cell(
        row=fila, column=1, value="Exceso de retención en origen por zona / mercado"
    ).font = _FONT_SECCION
    fila += 1
    ws.cell(
        row=fila, column=1,
        value=(
            "Agrupación orientativa (por zona, desglosada por mercado dentro de la zona EURO): "
            "puede no coincidir exactamente con el país de residencia fiscal real de la empresa. "
            "Comprueba siempre el certificado de retenciones de tu bróker."
        ),
    ).font = _FONT_NOTA
    fila += 1
    _escribir_cabecera(ws, fila, ["Zona / mercado", "Exceso sobre 15%"])
    fila_inicio_zonas = fila + 1
    fila += 1
    for item in resumen["exceso_origen_por_zona"]:
        ws.cell(row=fila, column=1, value=item["zona"]).font = _FONT_NORMAL
        celda = ws.cell(row=fila, column=2, value=item["importe"])
        celda.font = _FONT_NORMAL
        celda.number_format = FORMATO_EUR
        fila += 1
    if fila == fila_inicio_zonas:
        ws.cell(
            row=fila, column=1,
            value="Ninguna retención en origen supera el 15% del bruto cobrado este año.",
        ).font = _FONT_NOTA
        fila += 1
    fila += 1

    # --- Ventas (plusvalía FIFO) --------------------------------------------
    ws.cell(row=fila, column=1, value="Ventas (plusvalía por lotes FIFO)").font = _FONT_SECCION
    fila += 1
    cabecera_ventas = [
        "Fecha", "Cartera", "Ticker", "Empresa",
        "Nº títulos", "Importe venta", "Coste (FIFO)", "Plusvalía",
    ]
    _escribir_cabecera(ws, fila, cabecera_ventas)
    fila_inicio_ventas = fila + 1
    fila += 1
    for item in resumen["listado_ventas"]:
        ws.cell(row=fila, column=1, value=item["fecha_mostrar"]).font = _FONT_NORMAL
        ws.cell(row=fila, column=2, value=item["cartera"]).font = _FONT_NORMAL
        ws.cell(row=fila, column=3, value=item["ticker"]).font = _FONT_NORMAL
        ws.cell(row=fila, column=4, value=item["empresa"]).font = _FONT_NORMAL
        celda_titulos = ws.cell(row=fila, column=5, value=item["num_titulos"])
        celda_titulos.font = _FONT_NORMAL
        celda_titulos.number_format = FORMATO_TITULOS
        for col, clave in ((6, "importe_venta"), (7, "coste"), (8, "plusvalia")):
            celda = ws.cell(row=fila, column=col, value=item[clave])
            celda.font = _FONT_NORMAL
            celda.number_format = FORMATO_EUR
        fila += 1
    fila_fin_ventas = fila - 1
    if fila_fin_ventas >= fila_inicio_ventas:
        ws.cell(row=fila, column=4, value="Total").font = _FONT_TOTAL
        for col, letra in ((6, "F"), (7, "G"), (8, "H")):
            celda = ws.cell(
                row=fila, column=col,
                value=f"=SUM({letra}{fila_inicio_ventas}:{letra}{fila_fin_ventas})",
            )
            celda.font = _FONT_TOTAL
            celda.number_format = FORMATO_EUR
        fila += 1
    else:
        ws.cell(row=fila, column=1, value="(sin ventas este año)").font = _FONT_NOTA
        fila += 1

    _autoajustar_columnas(ws, [12, 16, 10, 24, 16, 14, 13, 15, 15])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
